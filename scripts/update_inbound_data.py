"""観光庁「宿泊旅行統計調査」第2次速報をJSONへ変換する。"""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import tempfile
import unicodedata
from dataclasses import dataclass
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook


SOURCE_PAGE = "https://www.mlit.go.jp/kankocho/tokei_hakusyo/shukuhakutokei.html"
ARCHIVE_PAGE = (
    "https://www.mlit.go.jp/kankocho/tokei_hakusyo/shukuhakutokei/"
    "kako_2jisokuho.html"
)
USER_AGENT = "UGATTA-InboundDataUpdater/1.0 (+https://ugatta-llc.com/)"
SURVEY_NOTE = "2026年1月から集計・層化基準変更"

PREFECTURES = [
    "北海道", "青森県", "岩手県", "宮城県", "秋田県", "山形県", "福島県",
    "茨城県", "栃木県", "群馬県", "埼玉県", "千葉県", "東京都", "神奈川県",
    "新潟県", "富山県", "石川県", "福井県", "山梨県", "長野県", "岐阜県",
    "静岡県", "愛知県", "三重県", "滋賀県", "京都府", "大阪府", "兵庫県",
    "奈良県", "和歌山県", "鳥取県", "島根県", "岡山県", "広島県", "山口県",
    "徳島県", "香川県", "愛媛県", "高知県", "福岡県", "佐賀県", "長崎県",
    "熊本県", "大分県", "宮崎県", "鹿児島県", "沖縄県",
]


@dataclass(frozen=True, order=True)
class Release:
    year: int
    month: int
    url: str
    label: str


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(unicodedata.normalize("NFKC", str(value)).split())


def parse_release(anchor_text: str, href: str, base_url: str) -> Release | None:
    text = normalize_text(anchor_text)
    if "第2次速報" not in text or ".xlsx" not in href.lower():
        return None
    match = re.search(r"(?P<year>20\d{2})年.*?(?P<month>\d{1,2})月分.*?第2次速報", text)
    if not match:
        return None
    return Release(
        year=int(match.group("year")),
        month=int(match.group("month")),
        url=urljoin(base_url, href.strip()),
        label=text,
    )


def fetch_page(session: requests.Session, url: str) -> tuple[str, BeautifulSoup]:
    response = session.get(url, timeout=45)
    response.raise_for_status()
    response.encoding = response.apparent_encoding or response.encoding
    return response.text, BeautifulSoup(response.content, "html.parser")


def discover_releases(session: requests.Session) -> tuple[list[Release], str]:
    main_text, main_soup = fetch_page(session, SOURCE_PAGE)
    pages: list[tuple[str, BeautifulSoup]] = [(SOURCE_PAGE, main_soup)]

    _, archive_soup = fetch_page(session, ARCHIVE_PAGE)
    year_pages: list[tuple[int, str]] = []
    for anchor in archive_soup.find_all("a", href=True):
        text = normalize_text(anchor.get_text(" ", strip=True))
        match = re.search(r"第2次速報値\((20\d{2})年\)", text)
        if match:
            year_pages.append(
                (int(match.group(1)), urljoin(ARCHIVE_PAGE, anchor["href"].strip()))
            )

    for _, page_url in sorted(year_pages, reverse=True)[:2]:
        _, soup = fetch_page(session, page_url)
        pages.append((page_url, soup))

    releases: dict[tuple[int, int], Release] = {}
    for page_url, soup in pages:
        for anchor in soup.find_all("a", href=True):
            release = parse_release(anchor.get_text(" ", strip=True), anchor["href"], page_url)
            if release:
                releases[(release.year, release.month)] = release

    if not releases:
        raise RuntimeError("観光庁ページから第2次速報Excelを特定できませんでした。")

    update_match = re.search(r"最終更新日\s*[:：]\s*(20\d{2})年(\d{1,2})月(\d{1,2})日", main_text)
    updated_at = (
        f"{int(update_match.group(1)):04d}-{int(update_match.group(2)):02d}-{int(update_match.group(3)):02d}"
        if update_match
        else date.today().isoformat()
    )
    return sorted(releases.values(), reverse=True), updated_at


def parse_numeric(value: Any) -> int:
    if value is None:
        return 0
    if isinstance(value, bool):
        raise ValueError("真偽値は統計値として扱えません。")
    if isinstance(value, (int, float)):
        if not math.isfinite(float(value)) or value < 0:
            raise ValueError(f"不正な統計値です: {value!r}")
        return int(round(value))
    text = normalize_text(value).replace("*", "").replace(",", "")
    if text in {"", "-", "―", "—"}:
        return 0
    if not re.fullmatch(r"\d+(?:\.0+)?", text):
        raise ValueError(f"数値として解析できません: {value!r}")
    return int(float(text))


def find_nationality_sheet(workbook: Any) -> Any:
    candidates = []
    for sheet in workbook.worksheets:
        preview = " ".join(
            normalize_text(cell)
            for row in sheet.iter_rows(min_row=1, max_row=min(8, sheet.max_row), values_only=True)
            for cell in row
            if cell is not None
        )
        if all(keyword in preview for keyword in ("施設所在地", "国籍", "外国人延べ宿泊者数")):
            if "市区町村" not in preview:
                candidates.append(sheet)
    if len(candidates) != 1:
        names = [sheet.title for sheet in candidates]
        raise RuntimeError(f"国籍別都道府県表を一意に特定できません: {names}")
    return candidates[0]


def clean_prefecture_label(value: Any) -> str:
    text = normalize_text(value).replace(" ", "")
    return re.sub(r"^\d{2}", "", text)


def extract_nationality_table(workbook: Any) -> dict[str, Any]:
    sheet = find_nationality_sheet(workbook)
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    header_index = next(
        (
            index
            for index, row in enumerate(rows)
            if {"韓国", "中国", "台湾"}.issubset({normalize_text(cell) for cell in row})
        ),
        None,
    )
    if header_index is None:
        raise RuntimeError("国籍別表の見出し行を特定できません。")

    header = rows[header_index]
    nationality_columns = [
        (index, normalize_text(value))
        for index, value in enumerate(header)
        if index >= 2 and normalize_text(value)
    ]
    if len(nationality_columns) < 10:
        raise RuntimeError("国籍列が不足しています。")

    title_text = " ".join(
        normalize_text(cell)
        for row in rows[:header_index]
        for cell in row
        if cell is not None
    )
    scope_match = re.search(r"((?:客室数|従業者数)[^()（）]*?以上の施設)", title_text)
    scope = scope_match.group(1) if scope_match else "国籍別集計の調査対象施設"

    national: dict[str, Any] | None = None
    prefectures: dict[str, dict[str, Any]] = {}
    for row in rows[header_index + 1 :]:
        label = normalize_text(row[0] if row else None)
        if not label:
            continue
        values = {
            name: parse_numeric(row[index] if index < len(row) else None)
            for index, name in nationality_columns
        }
        total = parse_numeric(row[1] if len(row) > 1 else None)
        cleaned = clean_prefecture_label(label)
        record = {"foreign_guest_nights": total, "nationality": values}
        if cleaned in PREFECTURES:
            prefectures[cleaned] = record
        elif national is None and re.search(r"(?:令和\d+年|20\d{2}年).*?\d{1,2}月", label):
            national = record

    missing = sorted(set(PREFECTURES) - set(prefectures))
    if missing:
        raise RuntimeError(f"都道府県データが不足しています: {missing}")
    if national is None:
        raise RuntimeError("全国値を特定できません。")
    return {
        "sheet_name": sheet.title,
        "survey_scope": scope,
        "national": national,
        "prefectures": prefectures,
    }


def load_release_table(session: requests.Session, release: Release) -> dict[str, Any]:
    response = session.get(release.url, timeout=90)
    response.raise_for_status()
    workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
    try:
        return extract_nationality_table(workbook)
    finally:
        workbook.close()


def ensure_finite(value: Any, path: str = "root") -> None:
    if isinstance(value, dict):
        for key, child in value.items():
            ensure_finite(child, f"{path}.{key}")
    elif isinstance(value, list):
        for index, child in enumerate(value):
            ensure_finite(child, f"{path}[{index}]")
    elif isinstance(value, float) and not math.isfinite(value):
        raise ValueError(f"NaNまたはInfinityを検出しました: {path}")


def validate_dataset(dataset: dict[str, Any], previous: dict[str, Any] | None = None) -> None:
    metadata = dataset.get("metadata", {})
    year = metadata.get("year")
    month = metadata.get("month")
    if not isinstance(year, int) or not 2007 <= year <= date.today().year + 1:
        raise ValueError("年月の年が不正です。")
    if not isinstance(month, int) or not 1 <= month <= 12:
        raise ValueError("年月の月が不正です。")
    if metadata.get("release_type") != "第2次速報":
        raise ValueError("第2次速報以外のデータは公開できません。")

    national = dataset.get("national", {})
    if not isinstance(national.get("foreign_guest_nights"), int) or national["foreign_guest_nights"] <= 0:
        raise ValueError("全国の外国人延べ宿泊者数が不正です。")
    if not national.get("nationality"):
        raise ValueError("全国の国籍データが空です。")
    if len(dataset.get("prefectures", {})) != 47 or set(dataset["prefectures"]) != set(PREFECTURES):
        raise ValueError("47都道府県が揃っていません。")

    nonzero_prefectures = 0
    for name, record in dataset["prefectures"].items():
        total = record.get("foreign_guest_nights")
        if not isinstance(total, int) or total < 0:
            raise ValueError(f"{name}の外国人延べ宿泊者数が不正です。")
        nonzero_prefectures += int(total > 0)
        nationalities = record.get("nationality")
        if not isinstance(nationalities, dict) or not nationalities:
            raise ValueError(f"{name}の国籍データが空です。")
        if any(not isinstance(value, int) or value < 0 for value in nationalities.values()):
            raise ValueError(f"{name}の国籍別数値が不正です。")
        monthly = record.get("monthly")
        if not isinstance(monthly, list) or not monthly:
            raise ValueError(f"{name}の月次データが空です。")
    if nonzero_prefectures < 42:
        raise ValueError("都道府県データに極端な欠落があります。")

    if previous:
        previous_total = previous.get("national", {}).get("foreign_guest_nights", 0)
        current_total = national["foreign_guest_nights"]
        if previous_total and current_total < previous_total * 0.2:
            raise ValueError("前回公開値と比べて全国値が極端に減少しています。")
    ensure_finite(dataset)


def build_dataset(
    session: requests.Session,
    releases: list[Release],
    updated_at: str,
    months: int,
) -> dict[str, Any]:
    selected = sorted(releases[:months], key=lambda item: (item.year, item.month))
    if len(selected) < min(months, 6):
        raise RuntimeError("月次推移を作成するための第2次速報が不足しています。")

    tables: dict[tuple[int, int], dict[str, Any]] = {}
    for release in selected:
        print(f"取得・解析: {release.year}年{release.month}月 第2次速報")
        tables[(release.year, release.month)] = load_release_table(session, release)

    latest = max(selected, key=lambda item: (item.year, item.month))
    latest_table = tables[(latest.year, latest.month)]
    monthly_national = []
    for release in selected:
        table = tables[(release.year, release.month)]
        monthly_national.append(
            {
                "year": release.year,
                "month": release.month,
                "foreign_guest_nights": table["national"]["foreign_guest_nights"],
                "survey_scope": table["survey_scope"],
            }
        )

    prefectures: dict[str, Any] = {}
    for prefecture in PREFECTURES:
        latest_record = latest_table["prefectures"][prefecture]
        monthly = []
        for release in selected:
            table = tables[(release.year, release.month)]
            monthly.append(
                {
                    "year": release.year,
                    "month": release.month,
                    "foreign_guest_nights": table["prefectures"][prefecture]["foreign_guest_nights"],
                    "survey_scope": table["survey_scope"],
                }
            )
        prefectures[prefecture] = {
            "foreign_guest_nights": latest_record["foreign_guest_nights"],
            "nationality": latest_record["nationality"],
            "monthly": monthly,
        }

    return {
        "metadata": {
            "source": "観光庁 宿泊旅行統計調査",
            "source_url": SOURCE_PAGE,
            "excel_url": latest.url,
            "year": latest.year,
            "month": latest.month,
            "release_type": "第2次速報",
            "updated_at": updated_at,
            "unit": "人泊",
            "survey_note": SURVEY_NOTE,
            "survey_scope": latest_table["survey_scope"],
            "nationality_coverage": "国籍（出身地）別集計表に掲載された国・地域",
            "sheet_name": latest_table["sheet_name"],
        },
        "national": {
            **latest_table["national"],
            "monthly": monthly_national,
        },
        "prefectures": prefectures,
    }


def serialize_json(payload: dict[str, Any]) -> str:
    return json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=False, allow_nan=False) + "\n"


def write_if_changed(path: Path, content: str) -> bool:
    if path.exists() and path.read_text(encoding="utf-8") == content:
        return False
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        "w", encoding="utf-8", dir=path.parent, prefix=f"{path.stem}.", suffix=".tmp.json", delete=False
    ) as handle:
        handle.write(content)
        temporary = Path(handle.name)
    try:
        json.loads(temporary.read_text(encoding="utf-8"))
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()
    return True


def ensure_github_actions_environment() -> None:
    """観光庁への取得処理をGitHub Actions内に限定する。"""
    if os.environ.get("GITHUB_ACTIONS") != "true":
        raise RuntimeError(
            "観光庁データの更新はGitHub Actionsからのみ実行できます。"
        )


def main() -> int:
    ensure_github_actions_environment()
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=Path("data/inbound/latest.json"))
    parser.add_argument("--metadata-output", type=Path, default=Path("data/inbound/metadata.json"))
    parser.add_argument("--months", type=int, default=12)
    args = parser.parse_args()

    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})
    releases, updated_at = discover_releases(session)
    latest = releases[0]
    print(f"最新第2次速報: {latest.year}年{latest.month}月 {latest.url}")

    previous = None
    if args.output.exists():
        previous = json.loads(args.output.read_text(encoding="utf-8"))
    dataset = build_dataset(session, releases, updated_at, args.months)
    validate_dataset(dataset, previous)

    metadata_document = {
        "metadata": dataset["metadata"],
        "available_prefectures": PREFECTURES,
        "available_months": [
            {"year": item["year"], "month": item["month"]}
            for item in dataset["national"]["monthly"]
        ],
    }
    changed_latest = write_if_changed(args.output, serialize_json(dataset))
    changed_metadata = write_if_changed(args.metadata_output, serialize_json(metadata_document))
    print("JSONを更新しました。" if changed_latest or changed_metadata else "新しい変更はありません。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
